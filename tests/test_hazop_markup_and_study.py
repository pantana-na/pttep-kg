"""Unit and Property-Based Tests for HAZOP P&ID Markup Ingestion, 3-Gate HITL, and 7-Tab Excel Export.

Governed by SPEC-20260831-HAZOP-MARKUP-INGESTION-AND-STUDY-LIFECYCLE.
"""

import os
import pytest
import openpyxl
from pathlib import Path
from hypothesis import given, settings, strategies as st
from fastapi.testclient import TestClient

from database.init_db import init_local_mock
from agents.hazop.agent import HazopStudyAgent
from agents.hazop.markup_parser import PidMarkupParser
from agents.hazop.ram_evaluator import (
    calculate_overall_severity,
    get_risk_rating,
    calculate_ipl_credit,
    evaluate_1st_risk,
    evaluate_2nd_risk,
    evaluate_deviation_risk
)
from agents.hazop.excel_exporter import export_hazop_study_to_excel
from server.main import app


@pytest.fixture
def mock_db():
    return init_local_mock("wiki")


@pytest.fixture
def hazop_agent(mock_db):
    return HazopStudyAgent(mock_db)


@pytest.fixture
def test_client():
    return TestClient(app)


# ==========================================
# 1. Unit Tests: Markup Parser & Hydration
# ==========================================

def test_markup_parser_node_23_02():
    """UT-MARKUP-01: Verifies parser correctly identifies Node 23-02 boundaries and equipment."""
    parser = PidMarkupParser()
    pdf_path = "hazop-example/Node 23-02.pdf"
    res = parser.extract_node_markup(pdf_path, filename="Node 23-02.pdf")

    assert res["node_id"] == "CDN-N02"
    assert res["colour_code"] == "Yellow"
    assert "E-2302A/B" in res["equipment_tags"]
    assert "E-2303" in res["equipment_tags"]
    assert "14780-8120-25-23-0005" in res["pid_drawings"]
    assert "Oxidate feed" in res["inlet_boundary"]


def test_markup_parser_node_23_03():
    """UT-MARKUP-02: Verifies parser correctly identifies Node 23-03 boundaries and equipment."""
    parser = PidMarkupParser()
    pdf_path = "hazop-example/Node 23-03.pdf"
    res = parser.extract_node_markup(pdf_path, filename="Node 23-03.pdf")

    assert res["node_id"] == "CDN-N03"
    assert res["colour_code"] == "Green"
    assert "E-2304" in res["equipment_tags"]
    assert "P-2301A/B" in res["equipment_tags"]


def test_node_hydration_and_confirmation(hazop_agent):
    """UT-CONFIRM-01: Verifies side-by-side operating parameters hydration and wiki file generation."""
    node_def = hazop_agent.parse_markup_and_hydrate("hazop-example/Node 23-02.pdf", filename="Node 23-02.pdf")
    assert len(node_def["parameters"]) >= 2
    assert any("E-2302A/B tube" in p["tag"] for p in node_def["parameters"])
    assert any("82–83 °C" in p["operating_condition"] for p in node_def["parameters"])

    conf_res = hazop_agent.confirm_node_definition(node_def)
    assert conf_res["status"] == "CONFIRMED"
    assert Path(conf_res["filepath"]).exists()


# ==========================================
# 2. Unit Tests: 3-Gate HITL Deviation Engine
# ==========================================

def test_hitl_gate_1_initial_risk(hazop_agent):
    """UT-HITL-01: Verifies Gate 1 unmitigated risk scoring with PEES severity maximum."""
    res = hazop_agent.evaluate_1st_risk_hitl(people=5, env=3, econ=4, social=3, initial_likelihood=4)
    assert res["overall_severity"] == 5
    assert res["initial_likelihood"] == 4
    assert res["initial_risk_rating"] == "Extreme"


def test_hitl_gate_2_safeguards_and_ipl(hazop_agent):
    """UT-HITL-02: Verifies Gate 2 safeguard proposal and SIL credit calculation."""
    safeguards = hazop_agent.propose_safeguards_hitl("E-2303", "Flow")
    assert len(safeguards) >= 2
    
    # Check IPL calculation
    sil1_credit = calculate_ipl_credit({"sil_rating": "SIL 1", "is_ipl": True})
    assert sil1_credit == 1
    sil2_credit = calculate_ipl_credit({"sil_rating": "SIL 2", "is_ipl": True})
    assert sil2_credit == 2
    non_ipl = calculate_ipl_credit({"description": "Operator procedure", "is_ipl": False})
    assert non_ipl == 0


def test_hitl_gate_3_mitigated_risk_and_recommendation(hazop_agent):
    """UT-HITL-03: Verifies Gate 3 mitigated risk calculation and recommendation trigger."""
    first = hazop_agent.evaluate_1st_risk_hitl(people=5, env=4, econ=5, social=4, initial_likelihood=4)
    safeguards = [
        {"description": "TXSHH-0502A/B (1oo2 SIL 1) trips UXV-0501/0502", "sil_rating": "SIL 1", "is_ipl": True},
        {"description": "FXSLL-0401A/B/C (2oo3 SIL 1) low feed trip", "sil_rating": "SIL 1", "is_ipl": True}
    ]
    res = hazop_agent.evaluate_2nd_risk_and_recommendation_hitl(
        deviation="Flow — No / Low Flow",
        cause="Feed valve fails closed",
        consequence="CHP thermal runaway",
        first_risk=first,
        confirmed_safeguards=safeguards
    )
    # L_init (4) - 2 IPL credits = L_mit (2)
    assert res["risk_assessment"]["mitigated_likelihood"] == 2
    assert res["risk_assessment"]["mitigated_risk_rating"] == "High"
    assert res["risk_assessment"]["requires_action"] is True
    assert "Verify" in res["ai_recommendation"] or len(res["ai_recommendation"]) > 5


# ==========================================
# 3. Unit Tests: 7-Tab Excel Export Equivalence
# ==========================================

def test_7_tab_excel_export_fidelity():
    """UT-EXCEL-01: Verifies exported Excel contains exactly 7 sheets with identical header structure."""
    study_meta = {
        "node_id": "CDN-N02",
        "name": "Preflash Column Feed-Heating Circuit",
        "markup_label": "Node 23-02",
        "unit": "CDN",
        "equipment_tags": ["E-2302A/B", "E-2303", "D-2308", "P-2308A/B"],
        "pid_drawings": ["14780-8120-25-23-0005", "14780-8120-25-23-0005A"]
    }
    rows = [
        {
            "ref": "1.1.1",
            "parameter": "Flow",
            "deviation": "No / Low Flow",
            "cause": "Feed pump trip",
            "consequence": "Heater dryout",
            "wo_l": 4, "wo_p": 5, "wo_en": 4, "wo_ec": 5, "wo_s": 4, "wo_rr": "Extreme",
            "safeguards": [{"description": "TXSHH-0501 (SIL 1)", "il_esd": "Yes", "ipl_credit": 1}],
            "w_l": 3, "w_rr": "High",
            "recommendation": "R-001: Install redundant interlock"
        }
    ]
    out_file = "output/exports/test_verification_export.xlsx"
    saved_path = export_hazop_study_to_excel(study_meta, rows, out_file)
    
    assert Path(saved_path).exists()
    wb = openpyxl.load_workbook(saved_path, data_only=True)
    expected_sheets = [
        "Cover Page", "HAZOP Information", "WorkSheet Index",
        "WorkSheet CDN-N02", "Action Items", "Risk Ranking", "Interlock-ESD Summary"
    ]
    assert wb.sheetnames == expected_sheets

    # Verify 27 columns on WorkSheet tab
    ws_node = wb["WorkSheet CDN-N02"]
    assert ws_node.max_column == 27


# ==========================================
# 4. Unit Tests: FastAPI Server Endpoints
# ==========================================

def test_api_upload_and_confirm(test_client):
    """UT-API-01: Verifies upload-markup and confirm-node endpoints."""
    pdf_bytes = b"%PDF-1.4 mock pdf with Node 23-02 annotations"
    
    resp_upload = test_client.post(
        "/api/v1/hazop/upload-markup",
        content=pdf_bytes,
        headers={"x-filename": "Node 23-02.pdf"}
    )
    assert resp_upload.status_code == 200
    data_upload = resp_upload.json()
    assert data_upload["node_definition"]["node_id"] == "CDN-N02"

    resp_conf = test_client.post("/api/v1/hazop/confirm-node", json=data_upload["node_definition"])
    assert resp_conf.status_code == 200
    assert resp_conf.json()["status"] == "CONFIRMED"


def test_api_3_gate_deviation(test_client):
    """UT-API-02: Verifies 3-gate HITL API workflow."""
    # Gate 1
    resp_g1 = test_client.post("/api/v1/hazop/deviation/1st-risk", json={
        "people": 5, "env": 4, "econ": 5, "social": 4, "initial_likelihood": 4
    })
    assert resp_g1.status_code == 200
    g1_data = resp_g1.json()["first_risk"]
    assert g1_data["initial_risk_rating"] == "Extreme"

    # Gate 2
    resp_g2 = test_client.post("/api/v1/hazop/deviation/safeguards", json={
        "equipment_tag": "E-2303", "deviation_type": "Flow"
    })
    assert resp_g2.status_code == 200
    g2_safeguards = resp_g2.json()["safeguards"]
    assert len(g2_safeguards) > 0

    # Gate 3
    resp_g3 = test_client.post("/api/v1/hazop/deviation/2nd-risk", json={
        "deviation": "Flow — No / Low Flow",
        "cause": "Loss of feed",
        "consequence": "CHP thermal runaway",
        "first_risk": g1_data,
        "confirmed_safeguards": g2_safeguards
    })
    assert resp_g3.status_code == 200
    g3_data = resp_g3.json()["result"]
    assert "mitigated_risk_rating" in g3_data["risk_assessment"]


def test_finalize_study_and_knowledge_catalog_sync(hazop_agent):
    """UT-FINAL-01: Verifies study finalization syncs Wiki, Spanner Graph, and Dataplex Knowledge Catalog."""
    meta = {
        "node_id": "CDN-N02",
        "name": "Preflash Column Feed-Heating Circuit",
        "unit": "CDN",
        "equipment_tags": ["E-2302A/B", "E-2303"],
        "pid_drawings": ["14780-8120-25-23-0005"]
    }
    rows = [{
        "ref": "1.1.1",
        "parameter": "Flow",
        "deviation": "No / Low Flow",
        "cause": "Feed pump trip",
        "consequence": "Heater dryout",
        "wo_rr": "Extreme",
        "w_rr": "High",
        "recommendation": "R-001: Install redundant interlock"
    }]
    
    res = hazop_agent.finalize_study_and_sync(meta, rows)
    assert res["status"] == "FINALIZED_AND_SYNCED"
    assert res["spanner_node_status"] == "COMPLETE"
    assert "phenol-psi" in res["knowledge_catalog_entry"]
    assert res["recommendations_registered"] == 1
    assert Path(res["excel_deliverable"]).exists()
    
    # Verify Knowledge Catalog entry in DB
    entry_id = res["knowledge_catalog_entry"].split("/")[-1]
    cat_entry = hazop_agent.db.get_knowledge_catalog_entry(entry_id)
    assert cat_entry is not None
    assert cat_entry["aspects"]["oems_005_process_safety_aspect"]["study_status"] == "COMPLETE"


def test_api_finalize_study(test_client):
    """UT-API-03: Verifies /api/v1/hazop/finalize-study endpoint."""
    resp = test_client.post("/api/v1/hazop/finalize-study", json={
        "study_metadata": {
            "node_id": "CDN-N02",
            "name": "Preflash Column Feed-Heating Circuit",
            "unit": "CDN",
            "equipment_tags": ["E-2303"],
            "pid_drawings": ["14780-8120-25-23-0005"]
        },
        "worksheet_rows": [
            {
                "ref": "1.1.1",
                "parameter": "Flow",
                "deviation": "No / Low Flow",
                "cause": "FCV-0501 fails closed",
                "consequence": "CHP thermal runaway",
                "wo_rr": "Extreme",
                "w_rr": "High",
                "recommendation": "R-001: Verify proof test interval"
            }
        ]
    })
    assert resp.status_code == 200
    data = resp.json()
    assert data["status"] == "FINALIZED_AND_SYNCED"
    assert data["spanner_node_status"] == "COMPLETE"
    assert "phenol-psi" in data["knowledge_catalog_entry"]


def test_discover_node_risks(hazop_agent):
    """UT-DISCOVER-01: Verifies all candidate risks are discovered at once for Node 23-02."""
    rows = hazop_agent.discover_node_risks("CDN-N02", ["E-2302A/B", "E-2303"])
    assert len(rows) >= 5
    assert any("Flow — No / Low Flow" in r["deviation"] for r in rows)
    assert any("Temperature — High Temperature" in r["deviation"] for r in rows)
    assert any("Pressure — High Pressure" in r["deviation"] for r in rows)
    assert any("Level — Low Level" in r["deviation"] for r in rows)
    
    # Verify candidate safeguards are populated
    for r in rows:
        assert len(r["available_safeguards"]) >= 1
        assert "wo_rr" in r
        assert "w_rr" in r


def test_api_discover_and_evaluate_row(test_client):
    """UT-API-DISCOVER-02: Verifies /api/v1/hazop/discover-risks and /api/v1/hazop/evaluate-row endpoints."""
    # 1. Discover risks
    resp_disc = test_client.post("/api/v1/hazop/discover-risks", json={
        "node_id": "CDN-N02",
        "equipment_tags": ["E-2302A/B", "E-2303"]
    })
    assert resp_disc.status_code == 200
    rows = resp_disc.json()["discovered_rows"]
    assert len(rows) >= 5

    # 2. In-row adjustment: Modify initial risk to P=5, En=4, Ec=5, S=4, L=4, and toggle safeguards
    target_row = rows[0]
    target_row["wo_p"] = 5
    target_row["wo_en"] = 4
    target_row["wo_ec"] = 5
    target_row["wo_s"] = 4
    target_row["wo_l"] = 4
    # Ensure 2 IPL credits selected
    target_row["available_safeguards"][0]["selected"] = True
    target_row["available_safeguards"][1]["selected"] = True

    resp_eval = test_client.post("/api/v1/hazop/evaluate-row", json=target_row)
    assert resp_eval.status_code == 200
    eval_row = resp_eval.json()["row"]

    # Initial severity max(5,4,5,4) = 5, L_wo=4 -> Extreme
    assert eval_row["wo_s_overall"] == 5
    assert eval_row["wo_l"] == 4
    assert eval_row["wo_rr"] == "Extreme"

    # Total active IPL = 2 credits -> Mitigated L = 4 - 2 = 2
    assert eval_row["total_ipl_credits"] == 2
    assert eval_row["w_l"] == 2
    # S=5, L=2 -> High
    assert eval_row["w_rr"] == "High"
    assert eval_row["recommendation"] != ""


# ==========================================
# 5. Property-Based Tests (Hypothesis)
# ==========================================

@given(
    st.integers(min_value=1, max_value=5),
    st.integers(min_value=1, max_value=5),
    st.integers(min_value=1, max_value=5)
)
def test_pbt_ram_monotonicity(severity, l1, l2):
    """PBT-RAM-MONO: For any fixed Severity, Risk(S, L1) <= Risk(S, L2) whenever L1 <= L2."""
    from agents.hazop.ram_evaluator import RISK_RANKS
    
    r1 = get_risk_rating(severity, min(l1, l2))
    r2 = get_risk_rating(severity, max(l1, l2))
    assert RISK_RANKS[r1] <= RISK_RANKS[r2]


@given(
    st.integers(min_value=1, max_value=5),
    st.integers(min_value=1, max_value=5),
    st.integers(min_value=1, max_value=5),
    st.integers(min_value=1, max_value=5),
    st.integers(min_value=1, max_value=5),
    st.integers(min_value=0, max_value=3)
)
def test_pbt_ipl_likelihood_bounds(p, en, ec, s, init_l, sil_credit):
    """PBT-IPL-BOUNDS: Mitigated likelihood must be strictly bounded in [1, 5]."""
    sg = [{"sil_rating": f"SIL {sil_credit}", "is_ipl": True}] if sil_credit > 0 else []
    first = evaluate_1st_risk(p, en, ec, s, init_l)
    second = evaluate_2nd_risk(first, sg)
    
    assert 1 <= second["mitigated_likelihood"] <= 5
    assert second["overall_severity"] == first["overall_severity"]


@settings(deadline=None)
@given(
    st.integers(min_value=1, max_value=5),
    st.integers(min_value=1, max_value=5),
    st.integers(min_value=1, max_value=5),
    st.integers(min_value=1, max_value=5),
    st.integers(min_value=1, max_value=5),
    st.booleans(),
    st.booleans()
)
def test_pbt_evaluate_row_monotonicity_with_safeguards(p, en, ec, s, init_l, sg1_sel, sg2_sel):
    """PBT-EVAL-ROW: Selecting additional valid IPL safeguards must never increase final risk rating."""
    from agents.hazop.ram_evaluator import RISK_RANKS
    mock_db = init_local_mock("wiki")
    agent = HazopStudyAgent(mock_db)

    row_data = {
        "ref": "1.1.1",
        "deviation": "Flow — No / Low Flow",
        "cause": "Pump trip",
        "consequence": "Overheating",
        "wo_p": p, "wo_en": en, "wo_ec": ec, "wo_s": s, "wo_l": init_l,
        "available_safeguards": [
            {"description": "SIS-1", "sil_rating": "SIL 1", "is_ipl": True, "ipl_credit": 1, "selected": sg1_sel},
            {"description": "SIS-2", "sil_rating": "SIL 2", "is_ipl": True, "ipl_credit": 2, "selected": sg2_sel}
        ]
    }
    eval_res = agent.evaluate_row(row_data)
    
    # Severity Invariance
    assert eval_res["w_s_overall"] == eval_res["wo_s_overall"]
    # Likelihood reduction bounded
    assert 1 <= eval_res["w_l"] <= eval_res["wo_l"]
    # Mitigated risk <= initial risk
    assert RISK_RANKS[eval_res["w_rr"]] <= RISK_RANKS[eval_res["wo_rr"]]


def test_generate_scenario_row_scenarios(hazop_agent):
    """UT-SCENARIO-01: Verifies AI scenario generation across diverse process failure modes."""
    # 1. Instrument air scenario
    row_air = hazop_agent.generate_scenario_row("CDN-N02", ["E-2302A/B", "E-2303"], "Loss of instrument air to steam control valve TV-0501", existing_rows_count=5)
    assert row_air["ref"] == "1.6.1"
    assert "Instrument Air" in row_air["deviation"]
    assert "TV-0501" in row_air["cause"]
    assert len(row_air["available_safeguards"]) >= 2
    assert row_air["wo_s_overall"] == 5
    assert row_air["w_l"] <= row_air["wo_l"]

    # 2. Power loss scenario
    row_power = hazop_agent.generate_scenario_row("CDN-N02", ["E-2302A/B", "E-2303"], "Power blackout to pump P-2308", existing_rows_count=6)
    assert row_power["ref"] == "1.7.1"
    assert "Power" in row_power["deviation"]

    # 3. Tube rupture scenario
    row_tube = hazop_agent.generate_scenario_row("CDN-N02", ["E-2302A/B", "E-2303"], "Tube rupture in E-2303", existing_rows_count=7)
    assert row_tube["ref"] == "1.8.1"
    assert "Tube Leak / Rupture" in row_tube["deviation"]

    # 4. Acid / metal contamination scenario
    row_chem = hazop_agent.generate_scenario_row("CDN-N02", ["E-2302A/B", "E-2303"], "Acid or transition metal contamination in oxidate feed", existing_rows_count=8)
    assert row_chem["ref"] == "1.9.1"
    assert "Chemical Contamination" in row_chem["deviation"]
    assert row_chem["wo_s_overall"] == 5


def test_api_generate_scenario_row(test_client):
    """UT-API-SCENARIO-02: Verifies /api/v1/hazop/generate-scenario-row endpoint."""
    resp = test_client.post("/api/v1/hazop/generate-scenario-row", json={
        "node_id": "CDN-N02",
        "equipment_tags": ["E-2302A/B", "E-2303"],
        "scenario_text": "Loss of instrument air supply to SC1.5 steam control valve TV-0501",
        "existing_rows_count": 5
    })
    assert resp.status_code == 200
    data = resp.json()
    assert data["status"] == "SUCCESS"
    row = data["row"]
    assert row["ref"] == "1.6.1"
    assert "Instrument Air" in row["deviation"]
    assert "wo_rr" in row
    assert "w_rr" in row
    assert len(row["available_safeguards"]) >= 2


@given(st.text(min_size=1, max_size=80))
def test_pbt_generate_scenario_row_invariants(scenario_text):
    """PBT-SCENARIO: Any arbitrary scenario text produces a strictly bounded, valid HAZOP row."""
    from agents.hazop.ram_evaluator import RISK_RANKS
    mock_db = init_local_mock("wiki")
    agent = HazopStudyAgent(mock_db)

    row = agent.generate_scenario_row("CDN-N02", ["E-2303"], scenario_text, existing_rows_count=5)
    
    assert row["ref"].startswith("1.6.")
    assert 1 <= row["wo_p"] <= 5
    assert 1 <= row["wo_en"] <= 5
    assert 1 <= row["wo_ec"] <= 5
    assert 1 <= row["wo_s"] <= 5
    assert 1 <= row["wo_l"] <= 5
    assert 1 <= row["w_l"] <= row["wo_l"]
    assert row["wo_s_overall"] == max(row["wo_p"], row["wo_en"], row["wo_ec"], row["wo_s"])
    assert RISK_RANKS[row["w_rr"]] <= RISK_RANKS[row["wo_rr"]]

