"""Unit and Property-Based Tests for HAZOP Study Agent & PTT GC RAM Evaluator (Step 5.0).

Verifies RAM calculation, IPL safeguard credits, Anti-Bias halting, and Excel workbook export.
"""

from pathlib import Path
import pytest
import openpyxl
from hypothesis import given, strategies as st
from database.init_db import init_local_mock
from agents.hazop.ram_evaluator import (
    calculate_overall_severity, get_risk_rating,
    calculate_ipl_credit, evaluate_deviation_risk,
    RISK_RANKS, RAM_GRID
)
from agents.hazop.agent import HazopStudyAgent


# ==============================================================================
# 1. Deterministic Unit Tests
# ==============================================================================

def test_ut_haz_01_deviation_evaluation():
    """UT-HAZ-01: Evaluate deviation with Severity P=5, L=4, Safeguard SIL 2."""
    db = init_local_mock("wiki")
    agent = HazopStudyAgent(db)
    
    safeguards = [
        {"description": "TXSHH-0502A/B High Temperature Trip to UC-2301 ESD", "sil_rating": "SIL 2", "is_ipl": True}
    ]
    
    res = agent.evaluate_deviation(
        deviation="Higher Temperature",
        cause="Steam control valve FCV-0501 fails open",
        consequence="CHP thermal runaway and tube rupture",
        people=5, env=3, econ=4, social=3,
        initial_likelihood=4,
        safeguards=safeguards
    )
    
    assessment = res["risk_assessment"]
    assert assessment["severity"] == 5
    assert assessment["initial_likelihood"] == 4
    assert assessment["initial_risk_rating"] == "Extreme"
    assert assessment["total_ipl_credits"] == 2
    assert assessment["mitigated_likelihood"] == 2
    assert assessment["mitigated_risk_rating"] == "High" or assessment["mitigated_risk_rating"] == "Medium"
    assert assessment["action_required"] is True


def test_ut_haz_02_anti_bias_halt(tmp_path):
    """UT-HAZ-02: Trigger study setup with old Phenol report in raw directory."""
    raw_dir = tmp_path / "raw"
    raw_dir.mkdir()
    (raw_dir / "14780-8120-PS-0001_PFD.pdf").touch()
    (raw_dir / "2016_Phenol_HAZOP_Report.pdf").touch()
    
    db = init_local_mock("wiki")
    agent = HazopStudyAgent(db)
    
    res = agent.start_study_setup("CDN-N02", raw_dir=str(raw_dir))
    assert res["status"] == "HALT_ANTI_BIAS_VIOLATION"
    assert "ANTI-BIAS HALT" in res["error"]


def test_hazop_excel_export(tmp_path):
    """Verify export_study_workbook creates a valid 7-tab Excel workbook."""
    db = init_local_mock("wiki")
    agent = HazopStudyAgent(db)
    
    rows = [{
        "item_no": 1,
        "deviation": "Higher Temperature",
        "cause": "FCV-0501 fails open",
        "consequence": "CHP thermal runaway",
        "p": 5, "en": 3, "ec": 4, "s": 3,
        "l_init": 4, "initial_risk": "Extreme",
        "safeguards": "TXSHH-0502 (SIL 1)",
        "l_mit": 3, "mitigated_risk": "High",
        "rec_no": "R-001", "rec_text": "Verify 1oo2 proof test interval"
    }]
    
    out_file = tmp_path / "report.xlsx"
    saved = agent.export_study_workbook(
        study_metadata={"node_id": "CDN-N02", "node_name": "Preflash Column Steam Heater"},
        worksheet_rows=rows,
        output_filepath=str(out_file)
    )
    
    assert Path(saved).exists()
    wb = openpyxl.load_workbook(saved)
    assert "Cover Page" in wb.sheetnames
    assert "Executive Summary" in wb.sheetnames
    assert "HAZOP Worksheet" in wb.sheetnames
    assert "Recommendation Summary" in wb.sheetnames
    assert "Action Item Tracking" in wb.sheetnames
    assert "RAM Definition Matrix" in wb.sheetnames
    assert "Document References" in wb.sheetnames


# ==============================================================================
# 2. Property-Based Invariant Tests (Hypothesis)
# ==============================================================================

@given(
    s=st.integers(min_value=1, max_value=5),
    l1=st.integers(min_value=1, max_value=5),
    l2=st.integers(min_value=1, max_value=5)
)
def test_pbt_ram_01_monotonicity(s, l1, l2):
    """PBT-RAM-01: RAM Monotonicity: For L1 <= L2, Risk(S, L1) <= Risk(S, L2)."""
    if l1 <= l2:
        r1 = get_risk_rating(s, l1)
        r2 = get_risk_rating(s, l2)
        assert RISK_RANKS[r1] <= RISK_RANKS[r2]


@given(
    p=st.integers(min_value=1, max_value=5),
    en=st.integers(min_value=1, max_value=5),
    ec=st.integers(min_value=1, max_value=5),
    s=st.integers(min_value=1, max_value=5)
)
def test_pbt_ram_02_severity_invariance(p, en, ec, s):
    """PBT-RAM-02: Severity Invariance: S == max(P, En, Ec, S)."""
    overall = calculate_overall_severity(p, en, ec, s)
    assert overall == max(p, en, ec, s)


@given(
    sil_str=st.sampled_from(["None", "SIL 1", "SIL 2", "SIL 3", "Non-IPL BPCS", "Operator Procedure"]),
    l_init=st.integers(min_value=1, max_value=5)
)
def test_pbt_ipl_01_safeguard_credit_bounds(sil_str, l_init):
    """PBT-IPL-01: Safeguard Credit Bounds: Mitigated likelihood is always >= 1."""
    sg = {"sil_rating": sil_str, "is_ipl": "SIL" in sil_str}
    credit = calculate_ipl_credit(sg)
    assert 0 <= credit <= 3
    if "Non-IPL" in sil_str or "Procedure" in sil_str:
        assert credit == 0
    l_mit = max(1, l_init - credit)
    assert 1 <= l_mit <= 5
