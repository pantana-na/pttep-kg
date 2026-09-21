"""Property-Based and Unit Tests for Neutral Refinery Data Sanitization.

Governed by SPEC-20260921-DATA-SANITIZATION-NEUTRAL-REFINERY.
Asserts zero leakage of proprietary entity identifiers (PTT, PTTGC, PTTEP, PPCL)
across all active codebase, database seeds, UI, and knowledge base markdown files,
while mathematically verifying RAM calculation and Excel export integrity.
"""

import os
import re
from pathlib import Path
import pytest
from hypothesis import given, strategies as st
import openpyxl

from app.hazop.ram_evaluator import get_risk_rating, calculate_ipl_credit, evaluate_deviation_risk
from app.hazop.excel_exporter import export_hazop_study_to_excel


# ==========================================
# 1. Unit Tests: Functional Integrity
# ==========================================

def test_ut_san_01_ram_lookups_preserved():
    """UT-SAN-01: Verifies RAM ratings remain identical after sanitization."""
    assert get_risk_rating(1, 1) == "Very Low"
    assert get_risk_rating(3, 3) == "Medium"
    assert get_risk_rating(4, 4) == "High"
    assert get_risk_rating(5, 4) == "Extreme"
    assert get_risk_rating(5, 5) == "Extreme"


def test_ut_san_02_excel_export_structure(tmp_path):
    """UT-SAN-02: Verifies Excel exporter creates all 7 tabs with 27 columns without error."""
    study_meta = {
        "node_id": "CDN-N02",
        "name": "Preflash Feed Circuit",
        "unit": "CDN",
        "markup_label": "Node 23-02",
        "equipment_tags": ["E-2302A/B", "V-2301"],
        "pid_drawings": ["14780-8120-25-23-0005"]
    }
    rows = [
        {
            "deviation": "High Pressure",
            "cause": "Control valve fail closed",
            "consequence": "Overpressurization of V-2301",
            "wo_l": 4, "wo_p": 4, "wo_en": 3, "wo_ec": 4, "wo_s": 3, "wo_rr": "High",
            "safeguards": [{"description": "PSV-2301 set at 12 bar", "il_esd": "No", "ipl_credit": 2}],
            "w_l": 2, "w_rr": "Medium",
            "recommendation": "Verify relief capacity"
        }
    ]
    out_file = str(tmp_path / "test_sanitized_export.xlsx")
    saved = export_hazop_study_to_excel(study_meta, rows, out_file)
    assert Path(saved).exists()

    wb = openpyxl.load_workbook(saved, data_only=True)
    expected_sheets = [
        "Cover Page", "HAZOP Information", "WorkSheet Index",
        "WorkSheet CDN-N02", "Action Items", "Risk Ranking", "Interlock-ESD Summary"
    ]
    assert wb.sheetnames == expected_sheets
    ws_node = wb["WorkSheet CDN-N02"]
    assert ws_node.max_column == 27


# ==========================================
# 2. Property-Based Tests: Safety Invariants
# ==========================================

DISALLOWED_PATTERN = re.compile(r"\b(PTT|PTTGC|PTTEP|PPCL)\b", re.IGNORECASE)

ACTIVE_SCAN_TARGETS = [
    Path("app"),
    Path("server"),
    Path("database/seeds"),
    Path("wiki")
]

ACTIVE_FILES = []
for target in ACTIVE_SCAN_TARGETS:
    if target.is_dir():
        for root, _, files in os.walk(target):
            for f in files:
                if f.endswith((".py", ".json", ".md", ".html", ".ts", ".tsx")):
                    ACTIVE_FILES.append(str(Path(root) / f))


@pytest.mark.parametrize("filepath", ACTIVE_FILES)
def test_pbt_san_01_zero_proprietary_identifiers_invariant(filepath):
    """PBT-SAN-01: Invariant verifying zero proprietary identifiers exist across active files."""
    with open(filepath, "r", encoding="utf-8", errors="ignore") as fp:
        content = fp.read()
    
    matches = DISALLOWED_PATTERN.findall(content)
    assert not matches, f"Disallowed identifiers {set(matches)} found in {filepath}"


@given(s=st.integers(min_value=1, max_value=5), l1=st.integers(min_value=1, max_value=4))
def test_pbt_san_02_ram_monotonicity(s, l1):
    """PBT-SAN-02: Verifies RAM risk severity ordering is monotonic for increasing likelihood."""
    severity_order = {"Very Low": 1, "Low": 2, "Medium": 3, "High": 4, "Extreme": 5}
    risk_1 = get_risk_rating(s, l1)
    risk_2 = get_risk_rating(s, l1 + 1)
    assert severity_order[risk_1] <= severity_order[risk_2]
